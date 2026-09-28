def solution(n):
    answer = []
    answer.append(n)
    while (n == 1):
        if (n % 2 == 0):
            n /= 2
            answer.append(n)
        else:
            n = 3 * n + 1
            answer.append(n)
    return answer



4주차 csv 파일 읽기까지 함


#---------------------------------------------------------------------------------------------------

import openpyxl
import pandas as pd
df_csv_exam = pd.read_csv("data/exam.csv")
df_csv_exam


----------------------------------------------------

df = pd.DataFrame({"name": ["김구름", "구름구름", "구우름", "구르미"],
                  "puppy": [10, 20, 30, 40],
                  "puppy_world": [20, 30, 40, 40]})
df.to_csv("output_sales.csv")





#---------------------------------------------------------------------------------------------------

import pandas as pd
dr_raw = pd.DataFrame({"groomi1" : [1, 2, 3],
                       "groomi2": [2, 3, 2]})
dr_raw = dr_raw.rename(columns = {"groomi2": "groomgroom"})
dr_raw

"""
	groomi1	groomgroom
0	1	2
1	2	3
2	3	2


"""

#---------------------------------------------------------------------------------------------------

head() 앞부분 출력
tail() 뒷부분 출력
shape 행, 열 개수 출력
info() 변수 속성 출력
describe() 요약 통계량 출력


-------------------------


sum() 내장 함수
.read_csv() 패키지 함수

.head() 매서드

#---------------------------------------------------------------------------------------------------


import openpyxl

wb = openpyxl.load_workbook('data/characters.xlsx')

print(wb.sheetnames)


sheet1 = wb['Sheet1']
sheet2 = wb['groomi_sheet']

sheet1.title = '구름이'
sheet2.title = '구름이22'
print(wb.sheetnames)
print(sheet1['A1'].value)
wb.create_sheet('groomi_sheet')
print(wb.sheetnames)
sheet2['B1'] = '김구름 멍멍이'
print(sheet2['B1'].value)
copysheet = wb.copy_worksheet(sheet2)
print(wb.sheetnames)

copysheet.title = 'groomgroom'
print(wb.sheetnames)
del wb['groomgroom']
print(wb.sheetnames)
sheet1['B1'].value = '구름구름'
print(sheet1['B1'].value)

# 💡 반드시 파일 저장 코드를 입력해야 실제 엑셀 파일에 반영됩니다!
wb.save('data/characters22.xlsx') # 파일 경로와 이름을 상황에 맞게 수정하세요.


#---------------------------------------------------------------------------------------------------


import csv
f = open('data/characters22.csv', 'r', encoding='cp949')
# f = open('data/characters22.csv', 'r', encoding='cp949', newline="")

rdr = csv.reader(f)
# rdr = csv.writer(f)

#rdr.writerow(['6', '김구름이', '흰색', '간식먹기', '곰돌이'])

for line in rdr:
    print(line)

f.close()


#---------------------------------------------------------------------------------------------------

import pandas as pd
groomi_friends = ["김구름", "이새싹", "박방울", "황참이"]
series_friends = pd.Series(groomi_friends)
series_friends

"""
0
0	김구름
1	이새싹
2	박방울
3	황참이
"""


#---------------------------------------------------------------------------------------------------


with open("groomi.txt", "a") as f:
    f.write("kimgroomgroom")

with open("groomi.txt", "r") as f2:
    data = f2.read()
    print(data)

with open('data/characters22.csv', 'a', encoding='cp949', newline='') as f:
    wr = csv.writer(f)
    
    wr.writerow(['030993920', '김김구름이', '흰색', '간식먹기', '곰돌이'])


#---------------------------------------------------------------------------------------------------

#파일 읽기 쓰기 등등
f = open("groomi.txt", "w")
f.write("kimgroomi mungmung\n")
f.write("kimgroomi mungmungg\n")
f.close()
f2 = open("groomi.txt", "r")
data = f2.read()
data2 = f2.readlines()
print(data)
for line in data2:
    print(line)
f2.close


#---------------------------------------------------------------------------------------------------

import seaborn
var = ["a", "a", "b", "c"]
seaborn.countplot(x = var)
#배열에서 수로 그래프 그려주는거

#---------------------------------------------------------------------------------------------------


import seaborn as sb
df = sb.load_dataset("titanic")
sb.countplot(data = df, x = "pclass")


#---------------------------------------------------------------------------------------------------


